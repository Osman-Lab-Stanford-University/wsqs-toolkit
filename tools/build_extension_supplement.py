"""Create a researcher-approved supplementary questionnaire from enabled EXT_ specs.

Only uses supplementary question registry. Never modifies CORE-8, survey masters,
or original scoring/model specifications. Emits editable DOCX and an empty CSV
header template, not any respondent-level data.
"""
from __future__ import annotations
import argparse
import csv
import json
import sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / 'src'))
from wsqs.extensions import read_registry


def generate(folder, output):
    folder = Path(folder).expanduser().resolve()
    output = Path(output).expanduser().resolve()
    if output.exists() or output == ROOT or ROOT in output.parents:
        raise ValueError('Output must be a NEW folder outside this software repository')
    instrument = json.loads((ROOT / 'spec/WSQS_Instrument_Schema.json').read_text(encoding='utf-8'))
    original = {x['Question_ID'] for x in instrument['questions']}
    with (ROOT / 'questionbook/feature_manifest.csv').open(newline='', encoding='utf-8') as f:
        protected = {row['feature'] for row in csv.DictReader(f)}
    active = read_registry(folder, original, protected)
    if not active:
        raise ValueError('No enabled supplementary research questions supplied')
    if len(active) > 31:
        raise ValueError('Excel EXTENSION ENTRY template supports up to 31 active questions')

    document = Document()
    section=document.sections[0]
    section.top_margin=Inches(.78)
    section.bottom_margin=Inches(.7)
    section.left_margin=Inches(.85)
    section.right_margin=Inches(.85)
    for style in document.styles:
        try:
            style.font.name='Times New Roman'
            style.font.color.rgb=RGBColor(0,0,0)
        except (AttributeError, TypeError):
            pass
    normal=document.styles['Normal']; normal.font.size=Pt(10.5)
    title=document.add_heading('WSQS | Supplementary research questions', level=0)
    title.alignment=WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph('Researcher-defined items only. Not part of the original CORE-8 score.')
    document.add_paragraph('Study-specific consent and ethics approval must be completed before administration.')
    document.add_paragraph('Nonidentifying record ID: __________________________')
    for index, item in enumerate(active, 1):
        header=document.add_paragraph(style='Normal')
        header.add_run(f'{index}. {item["id"]}  {item["question"]}').bold=True
        header.paragraph_format.keep_with_next=True
        if item['kind']=='single_select':
            for option in item['options']:
                paragraph=document.add_paragraph('□  '+option)
                paragraph.paragraph_format.left_indent=Inches(.2)
                paragraph.paragraph_format.space_after=Pt(2)
                paragraph.paragraph_format.keep_with_next=True
        else:
            p=document.add_paragraph('Response: __________________________________________')
            p.paragraph_format.keep_with_next=True
        p=document.add_paragraph('Applicability: '+item['applicability'])
        p.paragraph_format.space_after=Pt(9)
    document.add_paragraph('These supplementary items do not alter the original WSQS score or original 33-predictor model.')
    # Hard enforcement of all actual runs, including inherited headings.
    for para in document.paragraphs:
        for run in para.runs:
            run.font.name='Times New Roman';run.font.color.rgb=RGBColor(0,0,0)
    output.mkdir(parents=True)
    document.save(output/'WSQS_SUPPLEMENTARY_RESEARCH_QUESTIONS.docx')
    with (output/'EXTENSION_ENTRY_BLANK_TEMPLATE.csv').open('x',newline='',encoding='utf-8') as f:
        csv.writer(f).writerow(['Record ID']+[item['id'] for item in active])
    with (output/'EXTENSION_QUESTION_CODEBOOK.csv').open('x',newline='',encoding='utf-8') as f:
        fields=['id','question','kind','options','domain','applicability','source','provenance','contributor','consent_review_required','scoring_policy']
        writer=csv.DictWriter(f,fieldnames=fields)
        writer.writeheader()
        for item in active:
            writer.writerow({field: (' | '.join(item.get('options', [])) if field == 'options' else item.get(field, '')) for field in fields})
    (output/'SUPPLEMENT_README.txt').write_text(
        'This file is a supplementary research survey, not a CORE-8 modification.\n'
        'To use it in the Excel workbook, enter each EXT_ identifier in EXTENSION ENTRY row 4,\n'
        'and enter the corresponding nonidentifying Record ID and response in rows 5–104.\n'
        'Keep question JSONs available. Export using:\n'
        'python tools/extract_research_entry.py --input /private/completed.xlsx '
        '--extensions /private/enabled_specs --output /private/export_01\n'
        'The exported CSV contains private respondent data; do not share it.\n'
        'Save/export DOCX to PDF in Word when a PDF is needed. The consent text is not finalized.\n', encoding='utf-8')
    (output/'BUILD_REPORT.json').write_text(json.dumps({
        'enabled_extension_count':len(active),
        'enabled_question_ids':[item['id'] for item in active],
        'original_core8_modified':False,
        'respondent_data_exported':False,
        'publication_authorized':False
    },indent=2)+'\n',encoding='utf-8')
    print(f'PASS | Generated editable supplementary survey for {len(active)} validated EXT_ questions')

if __name__ == '__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--extensions',type=Path,required=True)
    parser.add_argument('--output',type=Path,required=True)
    args=parser.parse_args()
    generate(args.extensions,args.output)
